import base64
from io import BytesIO
from pathlib import Path
import sys
import unicodedata

from PIL import Image
from docx.document import Document as _DocumentClass
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _find_root():
    for base in (Path.cwd(), *Path.cwd().parents):
        if (base / "Public" / "Iconos").exists():
            return base
    return Path.cwd()


def _find_legend_dir():
    root = _find_root()
    candidates = (
        Path(__file__).resolve().parent / "serval_eider_legend",
        Path(sys.prefix) / "serval_eider_legend",
        root / "serval_eider_patch" / "legend_b64",
    )
    for path in candidates:
        if path.exists() and list(path.glob("part*.txt")):
            return path
    return candidates[-1]


_B64_DIR = _find_legend_dir()
_LEGEND_BYTES = None
_ORIGINAL_SAVE = None


def _legend_bytes():
    """Carga el arte enviado por Eider y lo convierte a PNG para Word."""
    global _LEGEND_BYTES
    if _LEGEND_BYTES is not None:
        return _LEGEND_BYTES

    parts = sorted(_B64_DIR.glob("part*.txt"))
    if not parts:
        raise FileNotFoundError("No se encontró la leyenda de alérgenos de Eider")

    # Los fragmentos pueden venir sin el padding '=' final o contener saltos de línea.
    # Se normalizan sin alterar el payload antes de una decodificación estricta.
    encoded = "".join("".join(p.read_text(encoding="ascii").split()) for p in parts)
    encoded += "=" * (-len(encoded) % 4)
    raw = base64.b64decode(encoded, validate=True)
    image = Image.open(BytesIO(raw))
    image.load()
    image = image.convert("RGB")
    out = BytesIO()
    image.save(out, format="PNG", optimize=True)
    _LEGEND_BYTES = out.getvalue()
    return _LEGEND_BYTES


def _normalize(text):
    text = unicodedata.normalize("NFD", text or "")
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return " ".join(text.upper().split())


def _remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _remove_old_body_legend(doc):
    """Elimina la leyenda antigua cuando alguna salida la generaba dentro del body."""
    for paragraph in list(doc.paragraphs):
        text = _normalize(paragraph.text)
        if (
            "GUIA DE ALERGENOS" in text
            or "INFORMACION ORIENTATIVA BASADA EN CARTA" in text
            or "DEBE VALIDARSE CON INGREDIENTES REALES" in text
        ):
            _remove_element(paragraph._element)

    terms = {
        "GLUTEN", "CRUSTACEOS", "HUEVOS", "PESCADO", "CACAHUETES", "SOJA", "LACTEOS",
        "FRUTOS DE CASCARA", "APIO", "MOSTAZA", "SESAMO", "SULFITOS", "ALTRAMUCES", "MOLUSCOS",
    }
    for table in list(doc.tables):
        joined = _normalize(" ".join(cell.text for row in table.rows for cell in row.cells))
        if len(table.rows) <= 3 and sum(term in joined for term in terms) >= 7:
            _remove_element(table._element)


def _clear(container):
    for child in list(container._element):
        container._element.remove(child)


def _set_cell_width(cell, width_cm):
    width = Cm(width_cm)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def _set_table_grid_widths(table, widths_cm):
    grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for idx, width_cm in enumerate(widths_cm):
        if idx < len(grid_cols):
            grid_cols[idx].set(qn("w:w"), str(int(Cm(width_cm).twips)))


def _set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")


def _add_picture(paragraph, width_cm):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.add_run().add_picture(BytesIO(_legend_bytes()), width=Cm(width_cm))


def _footer_variants(section):
    return section.footer, section.first_page_footer, section.even_page_footer


def _single_footer(section, usable_cm):
    # Pie real: la leyenda queda anclada abajo y fuera del flujo del contenido.
    if section.bottom_margin < Cm(3.15):
        section.bottom_margin = Cm(3.15)
    section.footer_distance = Cm(0.18)
    for footer in _footer_variants(section):
        footer.is_linked_to_previous = False
        _clear(footer)
        _add_picture(footer.add_paragraph(), min(18.0, usable_cm))


def _book_footer(section, usable_cm):
    # Modo libro: una leyenda al pie de cada mitad, sin invadir el cuerpo.
    if section.bottom_margin < Cm(2.55):
        section.bottom_margin = Cm(2.55)
    section.footer_distance = Cm(0.12)
    gutter_cm = 0.8
    side_cm = max(6.0, (usable_cm - gutter_cm) / 2.0)
    widths = (side_cm, gutter_cm, side_cm)

    for footer in _footer_variants(section):
        footer.is_linked_to_previous = False
        _clear(footer)
        table = footer.add_table(rows=1, cols=3, width=Cm(usable_cm))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _set_table_grid_widths(table, widths)
        for idx, width in enumerate(widths):
            cell = table.cell(0, idx)
            _set_cell_width(cell, width)
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for idx in (0, 2):
            cell = table.cell(0, idx)
            cell.text = ""
            _add_picture(cell.paragraphs[0], max(5.5, side_cm - 0.35))
        table.cell(0, 1).text = ""


def _prepare(doc):
    _remove_old_body_legend(doc)
    for section in doc.sections:
        usable_cm = max(
            1.0,
            (section.page_width - section.left_margin - section.right_margin) / 360000.0,
        )
        if section.page_width > section.page_height:
            _book_footer(section, usable_cm)
        else:
            _single_footer(section, usable_cm)


def _install():
    global _ORIGINAL_SAVE
    if getattr(_DocumentClass.save, "_eider_legend_patch", False):
        return
    _ORIGINAL_SAVE = _DocumentClass.save

    def save_with_eider_legend(self, path_or_stream):
        _prepare(self)
        return _ORIGINAL_SAVE(self, path_or_stream)

    save_with_eider_legend._eider_legend_patch = True
    _DocumentClass.save = save_with_eider_legend


_install()
