"""Local reproducible audit; no external API calls."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import json
from unittest.mock import patch
from io import BytesIO
from zipfile import ZipFile
from docx import Document
import qwen_core

MENU = {'restaurant_name': 'Restaurante QA', 'texto_extra': 'IVA incluido', 'categories': [
    {'name': 'Entrantes', 'category_text': 'Para compartir', 'dishes': [
        {'number': '1', 'name': 'Ensalada de queso y nueces', 'description': 'Tomate fresco', 'price': '12,50', 'allergens': ['lacteos', 'frutos de cascara'], 'review_notes': []},
        {'number': '2', 'name': 'Arroz con verduras', 'description': 'Sin gluten', 'price': '14', 'allergens': [], 'review_notes': []}]}]}

def load_app():
    with patch.object(qwen_core, 'ensure_configured'):
        import app
    return app

if __name__ == '__main__':
    app = load_app()
    out = Path(sys.argv[1] if len(sys.argv) > 1 else 'qa/before')
    out.mkdir(parents=True, exist_ok=True)
    report = {}
    for name, build in {
        'allergens': app.create_word,
        'no_allergens': app.create_client_word_without_allergens,
        'clean': app.create_clean_word,
        'editable': app.create_editable_word_clean_template,
        'premium': app.create_premium_editable_word,
        'book': app.create_landscape_book_word,
    }.items():
        try:
            buf = build(MENU)
            raw = buf.getvalue()
            (out / (name + '.docx')).write_bytes(raw)
            doc = Document(BytesIO(raw))
            report[name] = {'body_tables': len(doc.tables), 'footer_pictures': len(doc.sections[0].footer._element.xpath('.//pic:pic')), 'footer_tables': len(doc.sections[0].footer.tables)}
        except Exception as exc:
            report[name] = {'error': str(exc)}
    for flag in (True, False):
        pdf = app.create_client_pdf_bytes(MENU, with_allergens=flag)
        report['pdf_' + str(flag)] = bool(pdf)
        if pdf:
            (out / ('pdf_' + str(flag) + '.pdf')).write_bytes(pdf)
    source = [{'category_index': 0, 'dish_index': i, 'name': d['name']} for i, d in enumerate(MENU['categories'][0]['dishes'])]
    partial = {'items': [{'category_index': 0, 'dish_index': 0, 'allergens': ['lacteos'], 'confidence': 'alta'}]}
    with patch.object(qwen_core, 'chat', return_value=(json.dumps(partial), {})):
        try:
            result, model = qwen_core.allergen_reasoning_json('test', json.dumps(source))
            report['partial_allergens_accepted'] = len(result['items']) < len(source)
        except Exception:
            report['partial_allergens_accepted'] = False
    (out / 'results.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps(report, indent=2))
