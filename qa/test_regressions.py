import copy
import hashlib
import json
from io import BytesIO
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from qa.audit_runtime import MENU, load_app
import qwen_core as q
from word_legend import LEGEND_PATH, set_eider_footer
from docx import Document
from PIL import Image

app = load_app()


class RegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        Path('qa/after').mkdir(parents=True, exist_ok=True)

    def test_eur_prices_do_not_duplicate_currency(self):
        from ai_design import _format_price
        for formatter in (app.format_price, _format_price):
            self.assertEqual(formatter('12,50 EUR'), '12,50 €')
            self.assertEqual(formatter('12,50 eur'), '12,50 €')
            self.assertEqual(formatter('12,50 €'), '12,50 €')
            self.assertEqual(formatter(''), '')

    def test_exact_artwork_in_all_word_exports(self):
        image = LEGEND_PATH.read_bytes()
        for builder in (app.create_word, app.create_clean_word,
                        app.create_editable_word_clean_template,
                        app.create_premium_editable_word, app.create_landscape_book_word):
            with self.subTest(builder=builder.__name__):
                raw = builder(MENU).getvalue()
                doc = Document(BytesIO(raw))
                with ZipFile(BytesIO(raw)) as archive:
                    self.assertIn(image, [archive.read(n) for n in archive.namelist() if n.startswith('word/media/')])
                for section in doc.sections:
                    for footer in (section.footer, section.first_page_footer, section.even_page_footer):
                        self.assertEqual(len(footer._element.xpath('.//pic:pic')), 1)
                        ext = footer._element.xpath('.//wp:extent')[0]
                        w, h = int(ext.get('cx')), int(ext.get('cy'))
                        with Image.open(LEGEND_PATH) as im:
                            self.assertAlmostEqual(w/h, im.width/im.height, places=4)
                        self.assertGreater(section.bottom_margin, h + section.footer_distance)
                        self.assertLessEqual(w, section.page_width-section.left_margin-section.right_margin)
                    set_eider_footer(section)
                    self.assertEqual(len(section.footer._element.xpath('.//pic:pic')), 1)
                # No original legend linked from body or header.
                for part in [doc.part] + [s.header.part for s in doc.sections]:
                    self.assertNotIn(image, [r.target_part.blob for r in part.rels.values() if not r.is_external and r.reltype.endswith('/image')])

    def test_without_allergens_is_clean(self):
        for columns in (False, True):
            doc = Document(app.create_client_word_without_allergens(MENU, two_columns=columns))
            self.assertEqual(len(doc.inline_shapes), 0)
            self.assertTrue(all(not s.footer._element.xpath('.//pic:pic') for s in doc.sections))
            self.assertIn('12,50', '\n'.join(p.text for p in doc.paragraphs))

    def test_all_templates_and_zip(self):
        for theme in app.EDITABLE_WORD_THEMES:
            for columns in (False, True):
                doc = Document(app.create_editable_word_clean_template(MENU, theme, columns))
                self.assertEqual(len(doc.sections[0].footer._element.xpath('.//pic:pic')), 1)
                if columns:
                    self.assertEqual(len(doc.sections), 2)
                    self.assertNotEqual(doc.sections[0]._sectPr.xpath('./w:cols')[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num'), '2')
        with ZipFile(app.create_editable_templates_zip(MENU)) as archive:
            self.assertEqual(len(archive.namelist()), len(app.EDITABLE_WORD_THEMES))

    def test_qwen_normalization_preserves_data(self):
        data = copy.deepcopy(MENU)
        data.update(header_text='Cabecera', footer_text='Pie', notes=['Aviso'])
        data['categories'][0]['notes'] = 'Nota de categoria'
        normalized = q._normalize_menu_payload({'data': data})
        self.assertEqual(normalized, data)
        for empty in ([], {}, {'categories': []}):
            with self.assertRaises(ValueError):
                q._normalize_menu_payload(empty)

    def test_partial_duplicate_and_invalid_allergen_answers(self):
        source = [{'category_index': 0, 'dish_index': i} for i in range(2)]
        item = {'category_index': 0, 'dish_index': 0, 'allergens': []}
        for items in ([], [item], [item, item], [dict(item, dish_index='0')]):
            self.assertFalse(q._valid_allergen_result({'items': items}, source))
        complete = {'items': [item, dict(item, dish_index=1)]}
        self.assertTrue(q._valid_allergen_result(complete, source))
        with patch.object(q, 'chat', side_effect=[(json.dumps({'items': [item]}), {}), (json.dumps(complete), {})]):
            result, model = q.allergen_reasoning_json('system', json.dumps(source))
            self.assertEqual(result, complete)
            self.assertEqual(model, q.QWEN_MAX)

    def test_truncated_response_rejected(self):
        with self.assertRaises(ValueError):
            q._extract_text({'choices': [{'finish_reason': 'length', 'message': {'content': '{}'}}]})

    def test_translation_preserves_prices_and_allergens(self):
        payload = app._translation_payload(MENU)
        payload['categories'][0]['dishes'][0]['name'] = 'Cheese salad'
        with patch.object(app, 'translate_json', return_value=payload):
            translated = app.translate_menu_data(MENU, 'Ingles')
        for original, translated_dish in zip(MENU['categories'][0]['dishes'], translated['categories'][0]['dishes']):
            self.assertEqual(original['price'], translated_dish['price'])
            self.assertEqual(original['allergens'], translated_dish['allergens'])
        payload['categories'][0]['dishes'].append(payload['categories'][0]['dishes'][0])
        with patch.object(app, 'translate_json', return_value=payload), self.assertRaises(ValueError):
            app.translate_menu_data(MENU, 'Ingles')

    def test_exports_wait_for_allergen_review_acknowledgement(self):
        from streamlit.testing.v1 import AppTest
        pending = copy.deepcopy(MENU)
        pending['categories'][0]['dishes'][0]['review_notes'] = ['Revisar la ficha técnica de la salsa.']
        at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'), default_timeout=90).run()
        at.session_state.menu_data = pending
        at.run()
        self.assertFalse(at.exception)
        labels = [button.label for button in at.get('download_button')]
        self.assertTrue(any('SIN ALÉRGENOS' in label for label in labels))
        self.assertFalse(any('CON ALÉRGENOS' in label or 'plantilla Word' in label or 'modo libro' in label for label in labels))
        ack = next(box for box in at.checkbox if 'He revisado estos casos' in box.label)
        ack.set_value(True).run()
        self.assertFalse(at.exception)
        labels = [button.label for button in at.get('download_button')]
        self.assertTrue(any('WORD CON ALÉRGENOS' in label for label in labels))
        self.assertTrue(any('plantilla Word' in label for label in labels))
        self.assertTrue(any('MODO LIBRO' in label for label in labels))

    def test_low_confidence_allergen_result_becomes_review_note(self):
        data = copy.deepcopy(MENU)
        result = {'items': [
            {'category_index': 0, 'dish_index': 0, 'allergens': ['lacteos'], 'confidence': 'baja', 'reason': 'Receta ambigua.'},
            {'category_index': 0, 'dish_index': 1, 'allergens': [], 'confidence': 'alta', 'reason': 'Sin ingredientes declarados.'},
        ]}
        with patch.object(app, 'allergen_reasoning_json', return_value=(result, 'qwen3.8-max')):
            app.infer_allergens_with_best_model(data)
        dish = data['categories'][0]['dishes'][0]
        self.assertTrue(any('Confianza baja' in note for note in dish['review_notes']))
        self.assertIn('Receta ambigua.', dish['review_notes'][-1])
        self.assertEqual(app.menu_preflight(data)['review_count'], 1)

    def test_radar_and_universal_extractor_navigation(self):
        from streamlit.testing.v1 import AppTest
        at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'), default_timeout=90).run()
        at.radio[0].set_value('📡 Radar de Clientes').run()
        self.assertFalse(at.exception)
        at.text_input[0].set_value('Restaurante QA')
        at.text_input[1].set_value('Madrid')
        at.button[0].click().run()
        self.assertFalse(at.exception)
        markdown = '\n'.join(item.value for item in at.markdown)
        self.assertIn('Google Maps', markdown)
        self.assertIn('Instagram', markdown)
        self.assertIn('TripAdvisor', markdown)
        at.radio[0].set_value('📄 Extractor de Texto Universal').run()
        self.assertFalse(at.exception)
        self.assertEqual(at.title[0].value, 'Extractor de Texto Plano 📄➡️📝')
        uploader = next(widget for widget in at.get('file_uploader') if widget.label == 'Sube tu archivo')
        uploader.set_value(('carta_qa.txt', b'Ensalada QA\n12,50 EUR', 'text/plain'))
        at.run()
        next(button for button in at.button if button.label == '🔄 Extraer Todo el Texto').click().run()
        self.assertFalse(at.exception)
        self.assertTrue(any(button.label == '⬇️ Descargar Word con Texto Literal' for button in at.get('download_button')))
        exported = Document(BytesIO(at.session_state.universal_bytes))
        self.assertIn('Ensalada QA', '\n'.join(paragraph.text for paragraph in exported.paragraphs))

    def test_text_extractors_read_pdf_and_docx(self):
        import fitz
        source_pdf = BytesIO()
        pdf_doc = fitz.open()
        pdf_page = pdf_doc.new_page()
        pdf_page.insert_text((72, 72), 'Carta QA - Ensalada 12,50 EUR')
        source_pdf.write(pdf_doc.tobytes())
        source_pdf.seek(0)
        self.assertIn('Ensalada', app.extract_text_from_pdf(source_pdf))
        source_docx = Document()
        source_docx.add_paragraph('Carta QA · Ensalada 12,50 EUR')
        docx_buffer = BytesIO()
        source_docx.save(docx_buffer)
        docx_buffer.seek(0)
        self.assertIn('Ensalada', app.extract_text_from_docx(docx_buffer))

    def test_image_fallback_and_real_transport_contract(self):
        calls = []
        def post(payload, timeout=240):
            calls.append(payload)
            content = '[]' if len(calls) == 1 else json.dumps(MENU)
            return {'choices': [{'message': {'content': content}, 'finish_reason': 'stop'}]}
        with patch.object(q, '_post_json', side_effect=post):
            result = q.menu_json_from_image('Extract menu', Image.new('RGB', (100, 100)))
        self.assertEqual(result, MENU)
        self.assertEqual(len(calls), 2)
        self.assertEqual(calls[0]['messages'][0]['content'][1]['image_url'], calls[1]['messages'][0]['content'][1]['image_url'])

    def test_pdf_exports_and_long_menu(self):
        import fitz
        long_menu = copy.deepcopy(MENU)
        long_menu['categories'][0]['dishes'] *= 50
        for allergens in (False, True):
            pdf = app.create_client_pdf_bytes(long_menu, with_allergens=allergens)
            self.assertIsNotNone(pdf)
            with fitz.open(stream=pdf, filetype='pdf') as doc:
                self.assertGreater(len(doc), 1)
                text = ''.join(p.get_text() for p in doc)
                self.assertEqual(text.count('Tomate fresco'), 50)
                self.assertEqual('1169/2011' in text, allergens)
        Path('qa/after/long.docx').write_bytes(app.create_word(long_menu).getvalue())

    def test_design_pdf_without_remote_generation(self):
        from ai_design import create_ai_menu_pdf
        background = BytesIO()
        Image.new('RGB', (600, 850), 'white').save(background, format='PNG')
        pdf = create_ai_menu_pdf(MENU, background.getvalue(), app.ICON_MAP)
        self.assertTrue(pdf.startswith(b'%PDF'))
        Path('qa/after/design.pdf').write_bytes(pdf)

    def test_all_premium_visual_structures_render(self):
        import fitz
        from pdf_backend import render_pdf
        for variant in ('cards', 'soft', 'compact', 'mesa', 'table'):
            with self.subTest(variant=variant):
                html = app.create_premium_v6_html(MENU, theme_key='cafe', variant=variant, show_dish_icons=True)
                self.assertIn('Ensalada de queso y nueces', html)
                self.assertEqual(html.count('class="guide-item"'), 14)
                pdf = render_pdf(html, app.BASE_DIR)
                self.assertTrue(pdf.startswith(b'%PDF'))
                with fitz.open(stream=pdf, filetype='pdf') as document:
                    text = app.normalize_text(''.join(page.get_text() for page in document)).lower()
                    self.assertIn('ensalada de queso y nueces', text)

    def test_ui_navigation_edit_and_downloads(self):
        from streamlit.testing.v1 import AppTest
        at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'), default_timeout=90).run()
        self.assertFalse(at.exception)
        self.assertEqual(len(at.radio), 1)
        at.session_state.menu_data = copy.deepcopy(MENU)
        at.run()
        self.assertFalse(at.exception)
        self.assertEqual(len(at.tabs), 7)
        self.assertGreaterEqual(len(at.get('download_button')), 9)
        price = next(w for w in at.text_input if w.label == 'Precio')
        price.set_value('19,75').run()
        self.assertFalse(at.exception)
        self.assertEqual(at.session_state.menu_data['categories'][0]['dishes'][0]['price'], '19,75')


if __name__ == '__main__':
    unittest.main(verbosity=2)
