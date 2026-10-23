"""Embed Eider's artwork with its white matte removed in real Word footers."""
from pathlib import Path
import base64
from functools import lru_cache
from PIL import Image
from docx.shared import Cm, Pt

PUBLIC_DIR = Path(__file__).resolve().parent / 'Public'
ORIGINAL_LEGEND_PATH = PUBLIC_DIR / 'leyenda_Eider_recreada_colores_HD.png'
LEGEND_PATH = PUBLIC_DIR / 'leyenda_Eider_recreada_colores_HD_recortada.png'
LEGEND_WIDTH = Cm(19)


@lru_cache(maxsize=1)
def eider_legend_data_uri():
    return 'data:image/png;base64,' + base64.b64encode(LEGEND_PATH.read_bytes()).decode('ascii')


def set_eider_footer(section):
    # One legend per physical page, also in landscape/book mode.
    width = min(LEGEND_WIDTH, section.page_width - section.left_margin - section.right_margin)
    with Image.open(LEGEND_PATH) as artwork:
        height = int(width * artwork.height / artwork.width)
    section.footer_distance = Cm(0.45)
    # Reserve the cropped image, baseline, and breathing room below body text.
    # Never float the legend over the body.
    section.bottom_margin = max(section.bottom_margin or 0, height + Cm(1.05))
    for footer in (section.footer, section.first_page_footer, section.even_page_footer):
        footer.is_linked_to_previous = False
        for child in list(footer._element):
            footer._element.remove(child)
        paragraph = footer.add_paragraph()
        paragraph.alignment = 1
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.left_indent = Pt(0)
        fmt.right_indent = Pt(0)
        fmt.first_line_indent = Pt(0)
        fmt.line_spacing = 1
        run = paragraph.add_run()
        run.font.size = Pt(1)
        # Explicit proportional dimensions avoid asymmetric image DPI metadata.
        picture = run.add_picture(str(LEGEND_PATH), width=width, height=height)
        picture._inline.docPr.set('descr', 'Leyenda original de Eider: 14 alérgenos')
