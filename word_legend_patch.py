from __future__ import annotations

from io import BytesIO
from pathlib import Path
import unicodedata

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx.document import Document as _DocumentClass
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

_ROOT = Path(__file__).resolve().parent
_ICON_DIR = _ROOT / "Public" / "Iconos"

_ORDER = [
    "gluten", "crustaceos", "huevos", "pescado", "cacahuetes", "soja", "lacteos",
    "frutos de cascara", "apio", "mostaza", "sesamo", "sulfitos", "altramuces", "moluscos",
]
_ICON_FILES = {
    "gluten": "gluten.png",
    "crustaceos": "gambas.png",
    "huevos": "huevo.png",
    "pescado": "pescado.png",
    "cacahuetes": "cacahuetes.png",
    "soja": "soja.png",
    "lacteos": "lacteos.png",
    "frutos de cascara": "frutos_secos.png",
    "apio": "apio.png",
    "mostaza": "mostaza.png",
    "sesamo": "sesamo.png",
    "sulfitos": "sulfitos.png",
    "altramuces": "altramuces.png",
    "moluscos": "moluscos.png",
}
_LABELS = {
    "gluten": "GLUTEN",
    "crustaceos": "CRUSTÁCEOS",
    "huevos": "HUEVOS",
    "pescado": "PESCADO",
    "cacahuetes": "CACAHUETES",
    "soja": "SOJA",
    "lacteos": "LÁCTEOS",
    "frutos de cascara": "FRUTOS DE\nCÁSCARA",
    "apio": "APIO",
    "mostaza": "MOSTAZA",
    "sesamo": "GRANOS DE\nSÉSAMO",
    "sulfitos": "DIÓXIDO DE\nAZUFRE\nY SULFITOS",
    "altramuces": "ALTRAMUCES",
    "moluscos": "MOLUSCOS",
}
_NOTICE = (
    "Informamos de acuerdo con el Reglamento de la U.E 1169/2011, "
    "que nuestros productos contienen o pueden contener los siguientes alérgenos."
)
_LEGEND_BYTES: bytes | None = None
_ORIGINAL_SAVE = None


def _font(size: int):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        "DejaVuSans-Bold.ttf",
        "Arial Bold.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except Exception:
            continue
    return ImageFont.load_default()


def _fit_font(draw, text: str, max_width: int, start_size: int = 25, min_size: int = 14):
    for size in range(start_size, min_size - 1, -1):
        font = _font(size)
        bbox = draw.textbbox((0, 0), text, font=font)
        if bbox[2] - bbox[0] <= max_width:
            return font
    return _font(min_size)


def _legend_bytes() -> bytes:
    global _LEGEND_BYTES
    if _LEGEND_BYTES is not None:
        return _LEGEND_BYTES

    width, height = 1800, 280
    canvas = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(canvas)
    draw.rectangle((5, 5, width - 6, height - 6), outline="#3b1f12", width=5)

    header_font = _fit_font(draw, _NOTICE, width - 80)
    draw.text((40, 17), _NOTICE, fill="black", font=header_font)

    label_font = _font(16)
    slot = (width - 54) / len(_ORDER)
    icon_size = 88
    icon_top = 67
    text_top = 166
    resampling = getattr(getattr(Image, "Resampling", Image), "LANCZOS")

    for index, allergen in enumerate(_ORDER):
        center_x = 27 + slot * (index + 0.5)
        icon_path = _ICON_DIR / _ICON_FILES[allergen]
        if icon_path.exists():
            try:
                icon = Image.open(icon_path).convert("RGBA")
                icon = ImageOps.contain(icon, (icon_size, icon_size), method=resampling)
                x = int(center_x - icon.width / 2)
                y = int(icon_top + (icon_size - icon.height) / 2)
                canvas.paste(icon, (x, y), icon)
            except Exception:
                pass

        label = _LABELS[allergen]
        bbox = draw.multiline_textbbox((0, 0), label, font=label_font, spacing=0, align="center")
        text_width = bbox[2] - bbox[0]
        draw.multiline_text(
            (int(center_x - text_width / 2), text_top),
            label,
            fill="black",
            font=label_font,
            spacing=0,
            align="center",
        )

    output = BytesIO()
    canvas.save(output, format="PNG", optimize=True)
    _LEGEND_BYTES = output.getvalue()
    return _LEGEND_BYTES


def _normalize(text: str) -> str:
    text = unicodedata.normalize("NFD", text or "")
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return " ".join(text.upper().split())


def _remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _remove_old_body_legend(doc):
    for paragraph in list(doc.paragraphs):
        text = _normalize(paragraph.text)
        if (
            "GUIA DE ALERGENOS" in text
            or "INFORMACION ORIENTATIVA BASADA EN CARTA" in text
            or "DEBE VALIDARSE CON INGREDIENTES REALES" in text
        ):
            _remove_element(paragraph._element)

    allergen_terms = {
        "GLUTEN", "CRUSTACEOS", "HUEVOS", "PESCADO", "CACAHUETES", "SOJA", "LACTEOS",
        "FRUTOS DE CASCARA", "APIO", "MOSTAZA", "SESAMO", "SULFITOS", "ALTRAMUCES", "MOLUSCOS",
    }
    for table in list(doc.tables):
        joined = _normalize(" ".join(cell.text for row in table.rows for cell in row.cells))
        hits = sum(1 for term in allergen_terms if term in joined)
        if hits >= 7 and len(table.rows) <= 3:
            _remove_element(table._element)


def _clear_container(container):
    for child in list(container._element):
        container._element.remove(child)


def _set_cell_width(cell, width_cm: float):
    width = Cm(width_cm)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def _add_picture(paragraph, width_cm: float):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(BytesIO(_legend_bytes()), width=Cm(width_cm))


def _apply_single_footer(section, usable_cm: float):
    footer = section.footer
    footer.is_linked_to_previous = False
    footer.distance = Cm(0.18)
    _clear_container(footer)
    p = footer.add_paragraph()
    _add_picture(p, min(18.0, usable_cm))


def _apply_book_footer(section, usable_cm: float):
    footer = section.footer
    footer.is_linked_to_previous = False
    footer.distance = Cm(0.12)
    _clear_container(footer)

    gutter_cm = 0.8
    side_cm = max(6.0, (usable_cm - gutter_cm) / 2.0)
    table = footer.add_table(rows=1, cols=3, width=Cm(usable_cm))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [side_cm, gutter_cm, side_cm]
    for idx, width in enumerate(widths):
        _set_cell_width(table.cell(0, idx), width)
        table.cell(0, idx).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for idx in (0, 2):
        cell = table.cell(0, idx)
        cell.text = ""
        _add_picture(cell.paragraphs[0], max(5.5, side_cm - 0.15))
    table.cell(0, 1).text = ""


def _apply_new_footer(doc):
    for section in doc.sections:
        usable_cm = max(
            1.0,
            (section.page_width - section.left_margin - section.right_margin) / 360000.0,
        )
        is_landscape = section.page_width > section.page_height
        if is_landscape:
            _apply_book_footer(section, usable_cm)
        else:
            _apply_single_footer(section, usable_cm)


def _prepare(doc):
    _remove_old_body_legend(doc)
    _apply_new_footer(doc)


def install_patch():
    global _ORIGINAL_SAVE
    if getattr(_DocumentClass.save, "_eider_legend_patch", False):
        return

    _ORIGINAL_SAVE = _DocumentClass.save

    def save_with_eider_legend(self, path_or_stream):
        try:
            _prepare(self)
        except Exception:
            pass
        return _ORIGINAL_SAVE(self, path_or_stream)

    save_with_eider_legend._eider_legend_patch = True
    _DocumentClass.save = save_with_eider_legend
